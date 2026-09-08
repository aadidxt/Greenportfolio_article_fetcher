from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from app.database import Database
from app.exporter import export_csv, export_xlsx
from app.models import ArticleDraft, SearchCandidate
from app.pipeline import FetchPipeline


class FakeSearch:
    def __init__(self, candidates):
        self.candidates = candidates
        self.windows = []

    def search(self, since, until):
        self.windows.append((since, until))
        return list(self.candidates)


class FakeExtractor:
    def __init__(self, drafts):
        self.drafts = drafts

    def extract(self, candidate):
        return self.drafts.get(candidate.url)


class FakeSheets:
    def __init__(self, fail=False):
        self.rows = {}
        self.fail = fail

    def append_articles(self, articles):
        if self.fail:
            raise RuntimeError("Sheets unavailable")
        for item in articles:
            self.rows.setdefault(item.normalized_url, item)
        return [item.id for item in articles]


class FakeEmail:
    def __init__(self):
        self.calls = []

    def send(self, **kwargs):
        self.calls.append(kwargs)


class Clock:
    def __init__(self, *values):
        self.values = iter(values)

    def __call__(self):
        return next(self.values)


def build_pipeline(test_settings, database, search, extractor, sheets, emailer, clock, **changes):
    return FetchPipeline(
        settings=replace(test_settings, **changes),
        database=database,
        search_service=search,
        extractor=extractor,
        sheets=sheets,
        emailer=emailer,
        clock=clock,
    )


def sample_data():
    candidate = SearchCandidate(
        "https://example.com/markets/story?utm_source=rss",
        "Green Portfolio on disciplined investing",
        datetime(2026, 9, 7, 9, tzinfo=UTC),
        "Example Finance",
        "test",
    )
    draft = ArticleDraft(
        publisher="Example Finance",
        published_at=datetime(2026, 9, 7, 9, tzinfo=UTC),
        title="Green Portfolio on disciplined investing",
        description="Green Portfolio shared its view on disciplined equity research for long-term investors.",
        url="https://example.com/markets/story",
        normalized_url="https://example.com/markets/story",
        source="test",
    )
    return candidate, draft


def test_first_run_then_second_run_is_incremental_and_append_only(test_settings):
    candidate, draft = sample_data()
    database = Database(test_settings.database_path)
    search = FakeSearch([candidate])
    sheets = FakeSheets()
    emailer = FakeEmail()
    clock = Clock(
        datetime(2026, 9, 7, 10, tzinfo=UTC),
        datetime(2026, 9, 7, 10, 1, tzinfo=UTC),
        datetime(2026, 9, 8, 10, tzinfo=UTC),
        datetime(2026, 9, 8, 10, 1, tzinfo=UTC),
    )
    pipeline = build_pipeline(
        test_settings,
        database,
        search,
        FakeExtractor({candidate.url: draft}),
        sheets,
        emailer,
        clock,
    )

    first = pipeline.run(fetch_id="first")
    assert first["status"] == "success"
    assert first["new_articles"] == 1
    assert len(database.all_articles()) == 1
    assert len(sheets.rows) == 1
    assert database.get_state("last_successful_fetch_id") == "first"

    latest_xlsx = export_xlsx(database.articles_for_fetch("first"))
    workbook = load_workbook(filename=__import__("io").BytesIO(latest_xlsx))
    assert workbook.active.max_row == 2
    assert workbook.active["A2"].value == "Example Finance"
    assert "Green Portfolio" in export_csv(database.all_articles()).decode("utf-8-sig")

    second = pipeline.run(fetch_id="second")
    assert second["status"] == "success"
    assert second["new_articles"] == 0
    assert second["duplicates"] == 1
    assert len(database.all_articles()) == 1
    assert len(sheets.rows) == 1
    assert search.windows[0][0] == datetime(2026, 8, 8, 10, tzinfo=UTC)
    assert search.windows[1][0] == datetime(2026, 9, 5, 10, 1, tzinfo=UTC)


def test_sheet_failure_preserves_state_and_is_recoverable(test_settings):
    candidate, draft = sample_data()
    database = Database(test_settings.database_path)
    search = FakeSearch([candidate])
    sheets = FakeSheets(fail=True)
    clock = Clock(
        datetime(2026, 9, 7, 10, tzinfo=UTC),
        datetime(2026, 9, 7, 10, 1, tzinfo=UTC),
    )
    pipeline = build_pipeline(
        test_settings,
        database,
        search,
        FakeExtractor({candidate.url: draft}),
        sheets,
        FakeEmail(),
        clock,
    )

    with pytest.raises(RuntimeError, match="Sheets unavailable"):
        pipeline.run(fetch_id="failed")
    assert database.get_state("last_successful_fetch") is None
    assert database.get_fetch("failed")["status"] == "failed"
    assert len(database.unsynced_articles()) == 1
    assert len(database.all_articles()) == 1


def test_scheduled_run_emails_new_and_complete_workbooks(test_settings):
    candidate, draft = sample_data()
    database = Database(test_settings.database_path)
    emailer = FakeEmail()
    pipeline = build_pipeline(
        test_settings,
        database,
        FakeSearch([candidate]),
        FakeExtractor({candidate.url: draft}),
        FakeSheets(),
        emailer,
        Clock(
            datetime(2026, 9, 7, 3, 30, tzinfo=UTC),
            datetime(2026, 9, 7, 3, 31, tzinfo=UTC),
        ),
        email_enabled=True,
    )
    result = pipeline.run(run_type="scheduled", fetch_id="weekly")

    assert result["status"] == "success"
    assert result["email_status"] == "success"
    assert len(emailer.calls) == 1
    assert emailer.calls[0]["new_count"] == 1
    assert emailer.calls[0]["new_xlsx"].startswith(b"PK")
    assert emailer.calls[0]["all_xlsx"].startswith(b"PK")


def test_all_extractions_failing_does_not_advance_success_state(test_settings):
    candidate, _ = sample_data()
    database = Database(test_settings.database_path)
    pipeline = build_pipeline(
        test_settings,
        database,
        FakeSearch([candidate]),
        FakeExtractor({}),
        FakeSheets(),
        FakeEmail(),
        Clock(
            datetime(2026, 9, 7, 10, tzinfo=UTC),
            datetime(2026, 9, 7, 10, 1, tzinfo=UTC),
        ),
    )

    with pytest.raises(RuntimeError, match="Every discovered result failed"):
        pipeline.run(fetch_id="all-failed")
    assert database.get_state("last_successful_fetch") is None
    assert database.get_fetch("all-failed")["status"] == "failed"


def test_upgrade_repairs_legacy_false_success_timestamp(test_settings):
    database = Database(test_settings.database_path)
    started = datetime(2026, 9, 8, 6, 20, tzinfo=UTC)
    database.create_fetch("legacy-false-success", started, "manual")
    from app.models import FetchCounters

    database.finish_fetch(
        "legacy-false-success",
        status="success",
        ended_at=datetime(2026, 9, 8, 6, 21, tzinfo=UTC),
        counters=FetchCounters(search_results=22, failed_articles=22),
    )
    database.set_state("last_successful_fetch_id", "legacy-false-success")
    database.set_state("last_successful_fetch", "2026-09-08T06:21:00Z")

    repaired = Database(test_settings.database_path)

    repaired_fetch = repaired.get_fetch("legacy-false-success")
    assert repaired_fetch["status"] == "failed"
    assert repaired_fetch["sheets_status"] == "not_attempted"
    assert repaired.get_state("last_successful_fetch") is None
    assert repaired.get_state("last_successful_fetch_id") is None
