from __future__ import annotations

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ArticleRecord
from .sheets import SHEET_HEADERS


def _rows(articles: list[ArticleRecord]):
    for article in articles:
        yield [
            article.publisher,
            article.published_at.date().isoformat(),
            article.title,
            article.description,
            article.url,
        ]


def export_csv(articles: list[ArticleRecord]) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(SHEET_HEADERS)
    writer.writerows(_rows(articles))
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def export_xlsx(articles: list[ArticleRecord]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Articles"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:E{max(1, len(articles) + 1)}"

    for column, header in enumerate(SHEET_HEADERS, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003134")
        cell.alignment = Alignment(vertical="center")

    for row in _rows(articles):
        worksheet.append(row)

    widths = [24, 16, 56, 86, 58]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

